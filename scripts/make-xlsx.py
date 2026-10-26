#!/usr/bin/env python3
"""Builds art/IMAGE-LIST.xlsx from art/images.json (run after: node scripts/build-data.mjs)."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
L = json.load(open(os.path.join(root, 'art/images.json'), encoding='utf-8'))
NAMES = [('abcd', 'A-Z Stories'), ('akshar-gu', 'Gujarati Akshar'), ('akshar-hi', 'Hindi Akshar'),
         ('numbers', 'Numbers'), ('deities', 'Deities'), ('avatars', 'Dashavatar')]
F = 'Arial'
hdr_fill = PatternFill('solid', fgColor='17124A')
thin = Side(style='thin', color='DDDDDD')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
sm = wb.active; sm.title = 'Summary'
HEAD = ['#', 'Done?', 'Label', 'Save as (file name)', 'Reference first?', 'Prompt (copy into Google Flow / Gemini)']
WID = [6, 10, 34, 46, 30, 110]
yes = PatternFill('solid', fgColor='C6EFCE')

for key, title in NAMES:
    rows = [i for i in L if i['s'] == key]
    ws = wb.create_sheet(title)
    ws.append(HEAD)
    for c in range(1, len(HEAD) + 1):
        x = ws.cell(1, c); x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = hdr_fill
        x.alignment = Alignment(vertical='center', wrap_text=True); x.border = box
    for n, i in enumerate(rows, 1):
        ref = ('Use ' + i['r']) if i['r'] else ''
        ws.append([n, 'No', i['l'], i['f'], ref, i['p']])
    for r in range(2, len(rows) + 2):
        for c in range(1, len(HEAD) + 1):
            x = ws.cell(r, c); x.font = Font(name=F, size=10); x.border = box
            x.alignment = Alignment(vertical='top', wrap_text=(c in (3, 4, 5, 6)))
        ws.cell(r, 2).alignment = Alignment(horizontal='center', vertical='top')
        ws.row_dimensions[r].height = 62
    for c, w in enumerate(WID, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:F{len(rows) + 1}'
    dv = DataValidation(type='list', formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(dv); dv.add(f'B2:B{len(rows) + 1}')
    ws.conditional_formatting.add(f'B2:B{len(rows) + 1}', CellIsRule(operator='equal', formula=['"Yes"'], fill=yes))

# summary with live formulas
sm['A1'] = 'Sanatan app: image checklist'; sm['A1'].font = Font(name=F, bold=True, size=16)
sm['A2'] = 'Set "Done?" to Yes on each section sheet after the picture is made and uploaded in studio.html. Totals update automatically.'
sm['A2'].font = Font(name=F, italic=True, size=10, color='666666')
for c, h in enumerate(['Section', 'Total', 'Done', 'Left', '% done'], 1):
    x = sm.cell(4, c, h); x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = hdr_fill; x.border = box
for n, (key, title) in enumerate(NAMES):
    r = 5 + n; last = len([i for i in L if i['s'] == key]) + 1
    sm.cell(r, 1, title)
    sm.cell(r, 2, f"=COUNTA('{title}'!D2:D{last})")
    sm.cell(r, 3, f"=COUNTIF('{title}'!B2:B{last},\"Yes\")")
    sm.cell(r, 4, f'=B{r}-C{r}')
    sm.cell(r, 5, f'=IF(B{r}=0,0,C{r}/B{r})'); sm.cell(r, 5).number_format = '0%'
    for c in range(1, 6): sm.cell(r, c).font = Font(name=F); sm.cell(r, c).border = box
t = 5 + len(NAMES)
sm.cell(t, 1, 'Total')
for c, col in ((2, 'B'), (3, 'C'), (4, 'D')): sm.cell(t, c, f'=SUM({col}5:{col}{t - 1})')
sm.cell(t, 5, f'=IF(B{t}=0,0,C{t}/B{t})'); sm.cell(t, 5).number_format = '0%'
for c in range(1, 6):
    x = sm.cell(t, c); x.font = Font(name=F, bold=True); x.border = box; x.fill = PatternFill('solid', fgColor='FFF2CC')
for c, w in zip('ABCDE', [26, 10, 10, 10, 10]): sm.column_dimensions[c].width = w
sm['A14'] = 'How to work'; sm['A14'].font = Font(name=F, bold=True)
for k, line in enumerate(['1. For scene stories, make scene 1 first and attach it as the reference for scenes 2, 3, ...',
                          '2. Copy the prompt from column F, paste into Google Flow / Gemini, generate the picture (16:9).',
                          '3. Upload it in studio.html (it saves as webp with the right name). Mark Done = Yes here.',
                          '4. Download the ZIP from studio.html and send it back.']):
    sm.cell(15 + k, 1, line).font = Font(name=F, size=10)

out = os.path.join(root, 'art/IMAGE-LIST.xlsx')
wb.save(out); print('saved', out, len(L), 'pictures')
