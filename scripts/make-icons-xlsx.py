#!/usr/bin/env python3
"""Builds art/ICON-LIST.xlsx: every animated icon the app needs, with generation prompts."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

STYLE = ("3D animated cartoon app icon, glossy rounded Pixar-style render, soft rim light, rich saturated colours, "
         "warm Indian Sanatan Dharma theme (gold, saffron, deep indigo, lotus pink), single centred subject filling ~80% of a square frame, "
         "plain pure white background (for easy background removal), soft contact shadow, no text, no letters, no watermark, 1:1, 1024x1024.")
LOOP = ("Animate this icon as a seamless 2-second loop at 24 fps, camera locked, background stays plain white, only the subject moves. ")

# group, id, name, where used, final size, format, idle animation (loop), tap animation, subject for the picture prompt
I = [
("App identity","app-icon","App icon (installed app)","Phone home screen / PWA icon 192 and 512","512x512 + 192x192","PNG (static)","none (static)","none","a golden glowing Om symbol sitting in the centre of an open pink lotus, a small lit diya on each side, deep indigo rounded-square badge behind"),
("App identity","app-icon-maskable","App icon (maskable, safe zone)","Android adaptive icon","512x512 (subject inside the central 70%)","PNG (static)","none","none","the same golden Om on a lotus, but smaller, on a full-bleed deep indigo square so edges can be cropped"),
("App identity","favicon","Browser tab icon","Browser tab / bookmarks","32x32 and 64x64","PNG + ICO","none","none","a very simple bold golden Om on an indigo circle, readable at tiny size"),
("App identity","splash-loader","Opening / loading screen","While the app starts","512x512","Lottie or WebM loop","the Om glows and pulses softly while a ring of small diyas lights up one by one around it","none","a large glowing golden Om inside a ring of eight tiny diyas"),
("App identity","cover-logo","Cover logo","Top of the home screen","768x384","PNG or Lottie","soft golden shimmer sweeping across, lotus petals gently opening and closing","tap: a small burst of flower petals","an open lotus holding a glowing Om, with a peacock feather on each side"),
("Home sections","sec-akshar-gu","Gujarati letters","Home tile","512x512","Lottie / WebM loop","the Gujarati letter અ gently floats up and down with tiny sparkles around it","tap: letter does a happy bounce","a big chunky 3D golden Gujarati letter અ standing on a little lotus pad"),
("Home sections","sec-akshar-hi","Hindi letters","Home tile","512x512","Lottie / WebM loop","the Devanagari letter अ floats gently, sparkles twinkle","tap: bounce","a big chunky 3D golden Devanagari letter अ standing on a little lotus pad"),
("Home sections","sec-numbers","Numbers","Home tile","512x512","Lottie / WebM loop","the numbers 1, 2 and 3 pop up one after another like a counting game, then settle","tap: all three jump","three chunky 3D numbers 1 2 3 made of polished saffron and gold beside a small japa mala"),
("Home sections","sec-deities","Gods and Goddesses","Home tile","512x512","Lottie / WebM loop","a soft halo of light rotates behind the trishul and damru, tiny stars twinkle","tap: damru shakes with a little sparkle burst","a 3D trishul with a small damru and a lotus, glowing halo behind (no human or deity face)"),
("Home sections","sec-avatars","Dashavatar","Home tile","512x512","Lottie / WebM loop","the Sudarshan chakra slowly spins while a small golden fish and a conch circle around it","tap: chakra spins fast once","a glowing golden Sudarshan chakra with a conch shell and a small golden fish beside it"),
("Home sections","sec-abcd","Sanatan A to Z book","Home tile","512x512","Lottie / WebM loop","an open storybook, pages flutter gently and golden letters A B C float up from it","tap: pages flip once","an open colourful storybook with floating 3D letters A B C and a small diya on top"),
("Home sections","sec-sanskrit","Sanskrit (coming soon)","Home tile","512x512","Lottie / WebM loop","a lotus bud slowly opens a little and closes again","tap: tiny wobble","a closed pink lotus bud over an old palm-leaf manuscript with a golden Om"),
("Header and controls","ui-home","Home","Top bar button","128x128","Lottie / SVG","tiny diya flame flickers on the temple top","tap: temple does a small hop","a cute small temple with a saffron flag and a diya at the door"),
("Header and controls","ui-language","Language","Top bar button","128x128","Lottie / SVG","the globe slowly turns and the three letters ગ, अ, A swap places around it","tap: spins once","a small 3D globe with three tiny floating letters, ગ अ A, orbiting around it"),
("Header and controls","ui-install","Add app to phone","Top bar button (only when installable)","128x128","Lottie / SVG","a down arrow gently bobs into the phone, a small diya glows on the phone screen","tap: phone does a happy shake","a phone showing a diya on its screen with a downward golden arrow above it"),
("Header and controls","ui-sound-on","Sound on / auto-read on","Top bar button","128x128","Lottie / SVG","sound waves pulse out of the conch","tap: conch bounce","a white conch shell blowing three curved golden sound waves"),
("Header and controls","ui-sound-off","Sound off","Top bar button","128x128","Lottie / SVG","a gentle breathing, no waves","tap: tiny nod","a white conch shell with a small crossed-out wave sign, calm and quiet"),
("Header and controls","ui-index","All letters (grid)","Top bar button","128x128","Lottie / SVG","tiles take turns lighting up","tap: tiles shuffle","nine small colourful 3D letter tiles in a 3x3 grid"),
("Header and controls","ui-prev","Previous page","Bottom navigation","160x160","Lottie / SVG","arrow sways slightly to the left","tap: quick push to the left with a petal puff","a golden arrow made of a lotus petal pointing left"),
("Header and controls","ui-next","Next page","Bottom navigation","160x160","Lottie / SVG","arrow sways slightly to the right","tap: quick push to the right with a petal puff","a golden arrow made of a lotus petal pointing right"),
("Header and controls","ui-listen","Listen / read aloud","Main button on every page","160x160","Lottie / SVG","sound waves ripple, tiny musical notes float up","tap: conch bounce","a bright blue conch shell with music notes and sound waves"),
("Header and controls","ui-pause","Pause / stop","Main button while speaking","160x160","Lottie / SVG","two bars breathe softly","tap: squash and stretch","two chunky rounded pause bars on a small glossy lotus"),
("Header and controls","ui-quiz","Quiz","Page button","160x160","Lottie / SVG","a question mark bounces while a small lightbulb flickers","tap: question mark flips to a star","a chunky 3D question mark with a tiny glowing lightbulb"),
("Header and controls","ui-star-off","I know this (not yet)","Page button","160x160","Lottie / SVG","outline star twinkles faintly","tap: fills with gold (see star-on)","a hollow translucent golden star with a faint glow"),
("Header and controls","ui-star-on","I know this (done)","Page button","160x160","Lottie / SVG","gold star shines and twinkles","tap: star pops with sparkles and a tiny Om appears","a solid glossy golden star with an Om engraved in the middle"),
("Header and controls","ui-zoom","Zoom picture","On every picture","96x96","Lottie / SVG","glass shimmers","tap: lens grows slightly","a magnifying glass with a golden frame and a small lotus in the lens"),
("Header and controls","ui-close","Close","Pop-ups","96x96","Lottie / SVG","none","tap: petals fold closed","a round pink lotus-petal button with a white cross"),
("Header and controls","tab-vowels","Vowels tab","Letters screen","128x128","Lottie / SVG","lotus petals open and close slowly","tap: bloom","a small open lotus with five petals"),
("Header and controls","tab-consonants","Consonants tab","Letters screen","128x128","Lottie / SVG","conch rocks slightly","tap: bounce","a small conch shell with a golden rim"),
("Feedback","fb-right","Correct answer","Quiz result","512x512","Lottie / WebM (play once)","none","a green tick grows with a burst of flower petals and sparkles, then settles","a chunky glossy green tick mark with flower petals and sparkles around it"),
("Feedback","fb-wrong","Try again","Quiz result","512x512","Lottie / WebM (play once)","none","a gentle friendly shake, a small peacock feather floats down, no sad faces","a soft orange circle with a friendly curved arrow and a peacock feather"),
("Feedback","fb-celebrate","All done celebration","End screen","768x768","Lottie / WebM (play once)","none","diyas light up in a row, fireworks of flowers burst, golden stars rain down","a big golden trophy-like diya plate overflowing with flowers and stars"),
("Feedback","confetti-diya","Confetti piece: diya","Falling confetti","128x128","PNG sprite","drifts and rotates while falling","none","one lit clay diya"),
("Feedback","confetti-lotus","Confetti piece: lotus","Falling confetti","128x128","PNG sprite","drifts and rotates while falling","none","one pink lotus flower"),
("Feedback","confetti-feather","Confetti piece: peacock feather","Falling confetti","128x128","PNG sprite","drifts and rotates while falling","none","one peacock feather"),
("Feedback","confetti-star","Confetti piece: star","Falling confetti","128x128","PNG sprite","drifts and rotates while falling","none","one golden star"),
("Feedback","confetti-marigold","Confetti piece: marigold","Falling confetti","128x128","PNG sprite","drifts and rotates while falling","none","one orange marigold flower"),
("Content chips","chip-points","Remember (key points)","Story page","96x96","SVG / Lottie","tiny sparkle","none","a golden star with a small sparkle"),
("Content chips","chip-words","New words","Story page","96x96","SVG / Lottie","page corner flutters","none","a small open book with a golden bookmark"),
("Content chips","chip-mantra","Mantra","Deity page","128x128","Lottie / SVG","Om glows softly","none","a golden Om inside a small circle of light"),
("Content chips","chip-vahan","Vehicle (vahan)","Deity page","96x96","SVG / Lottie","wheel turns","none","a small golden chariot wheel"),
("Content chips","chip-loves","Favourite offering","Deity page","96x96","SVG / Lottie","petals sway","none","a small plate with a flower and a sweet laddu"),
("Content chips","chip-yuga","Yuga (age)","Avatar page","96x96","SVG / Lottie","the wheel of time slowly turns","none","a golden wheel with a small sun and moon"),
("Content chips","chip-form","Form","Avatar page","96x96","SVG / Lottie","soft shimmer","none","a small glowing lotus with a mirror-like sparkle"),
("Rewards","reward-gu","Badge: Gujarati letters done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with the letter અ and a lotus ribbon"),
("Rewards","reward-hi","Badge: Hindi letters done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with the letter अ and a lotus ribbon"),
("Rewards","reward-numbers","Badge: Numbers done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with the number 108 and a lotus ribbon"),
("Rewards","reward-deities","Badge: Gods and Goddesses done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with a trishul and damru and a lotus ribbon"),
("Rewards","reward-avatars","Badge: Dashavatar done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with the Sudarshan chakra and a lotus ribbon"),
("Rewards","reward-abcd","Badge: A to Z done","Rewards","512x512","Lottie / PNG","medal shines, ribbon sways","pop with sparkles","a round golden medal with letters A to Z around a lotus and a lotus ribbon"),
("Rewards","reward-streak","Daily streak","Home screen","256x256","Lottie / PNG","the diya flame flickers brighter as the streak grows","flame leaps","a big glowing diya with a tall friendly flame"),
("Empty and system","sys-loading","Loading","Anywhere something loads","256x256","Lottie / WebM loop","the Sudarshan chakra spins smoothly","none","a small spinning golden chakra"),
("Empty and system","sys-offline","Offline","When there is no internet","256x256","Lottie / PNG","a cloud with a diya glows softly","none","a calm cloud holding a small lit diya, with a tiny crossed wifi sign"),
("Empty and system","sys-empty","Nothing here yet","Empty lists","256x256","Lottie / PNG","a lotus bud sways","none","a closed lotus bud on a small leaf"),
]

F='Arial'; thin=Side(style='thin',color='DDDDDD'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
hf=PatternFill('solid',fgColor='17124A')
wb=Workbook(); sm=wb.active; sm.title='Summary'; ws=wb.create_sheet('Icons')
H=['#','Done?','Group','Icon file name','Icon','Used where','Final size','Format','Idle loop animation','Tap animation','Step 1: picture prompt (Gemini / Flow)','Step 2: animation prompt (Flow / video)']
W=[5,9,18,22,28,28,18,18,40,40,70,60]
ws.append(H)
for c in range(1,len(H)+1):
    x=ws.cell(1,c); x.font=Font(name=F,bold=True,color='FFFFFF'); x.fill=hf; x.border=box; x.alignment=Alignment(vertical='center',wrap_text=True)
for n,(g,i,name,where,size,fmt,idle,tap,subj) in enumerate(I,1):
    anim=(LOOP+idle) if idle.lower() not in ('none','none (static)') else ''
    if tap.lower()!='none': anim+=(' ' if anim else '')+'Separate one-shot version: '+tap
    ws.append([n,'No',g,f'art/icons/{i}',name,where,size,fmt,idle,tap,f'{STYLE} Subject: {subj}.',anim or 'Static icon: no animation needed.'])
for r in range(2,len(I)+2):
    for c in range(1,len(H)+1):
        x=ws.cell(r,c); x.font=Font(name=F,size=10); x.border=box; x.alignment=Alignment(vertical='top',wrap_text=True)
    ws.cell(r,2).alignment=Alignment(horizontal='center',vertical='top'); ws.row_dimensions[r].height=95
for c,w in enumerate(W,1): ws.column_dimensions[get_column_letter(c)].width=w
ws.freeze_panes='E2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(H))}{len(I)+1}'
dv=DataValidation(type='list',formula1='"Yes,No"'); ws.add_data_validation(dv); dv.add(f'B2:B{len(I)+1}')
ws.conditional_formatting.add(f'B2:B{len(I)+1}',CellIsRule(operator='equal',formula=['"Yes"'],fill=PatternFill('solid',fgColor='C6EFCE')))

sm['A1']='Sanatan app: animated icon checklist'; sm['A1'].font=Font(name=F,bold=True,size=16)
for c,h in enumerate(['Group','Total','Done','Left','% done'],1):
    x=sm.cell(3,c,h); x.font=Font(name=F,bold=True,color='FFFFFF'); x.fill=hf; x.border=box
groups=[]
for g,*_ in I:
    if g not in groups: groups.append(g)
last=len(I)+1
for n,g in enumerate(groups):
    r=4+n
    sm.cell(r,1,g); sm.cell(r,2,f'=COUNTIF(Icons!$C$2:$C${last},A{r})'); sm.cell(r,3,f'=COUNTIFS(Icons!$C$2:$C${last},A{r},Icons!$B$2:$B${last},"Yes")')
    sm.cell(r,4,f'=B{r}-C{r}'); sm.cell(r,5,f'=IF(B{r}=0,0,C{r}/B{r})'); sm.cell(r,5).number_format='0%'
    for c in range(1,6): sm.cell(r,c).font=Font(name=F); sm.cell(r,c).border=box
t=4+len(groups); sm.cell(t,1,'Total')
for c,col in ((2,'B'),(3,'C'),(4,'D')): sm.cell(t,c,f'=SUM({col}4:{col}{t-1})')
sm.cell(t,5,f'=IF(B{t}=0,0,C{t}/B{t})'); sm.cell(t,5).number_format='0%'
for c in range(1,6):
    x=sm.cell(t,c); x.font=Font(name=F,bold=True); x.border=box; x.fill=PatternFill('solid',fgColor='FFF2CC')
for c,w in zip('ABCDE',[24,10,10,10,10]): sm.column_dimensions[c].width=w
notes=['How to make each icon:',
 '1. Copy "Step 1" prompt into Gemini / Google Flow and make a still picture. Keep the same style for all icons.',
 '2. Remove the white background (Canva or Photoshop "remove background") so the icon is transparent.',
 '3. If column I/J says there is an animation, give the picture to Flow/Veo with the "Step 2" prompt (2 second loop).',
 '4. Save as the file name in column D (png for stills, webm or lottie json for animations) and mark Done = Yes.',
 'Tip: UI chrome icons (arrows, close, zoom, pause) can also be drawn as light animated SVG by the developer. Ask if you want that instead of AI pictures.']
for k,l in enumerate(notes): sm.cell(t+2+k,1,l).font=Font(name=F,size=10,bold=(k==0))
out=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),'art/ICON-LIST.xlsx'); wb.save(out); print('saved',len(I),'icons')
